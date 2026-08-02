"""Парсер Excel (.xlsx). Файл (bytes) → список сырых строк (заголовок→значение)."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.modules.import_data.parsers.base import SOURCE_ROW_NUMBER


def parse_excel(content: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header = next(rows_iter)
    except StopIteration:
        return []

    headers = [str(h).strip() if h is not None else "" for h in header]
    raw_rows: list[dict] = []
    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None for c in row):
            continue
        raw = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
        raw[SOURCE_ROW_NUMBER] = row_number
        raw_rows.append(raw)
    wb.close()
    return raw_rows
